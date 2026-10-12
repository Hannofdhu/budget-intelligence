"""Read original OOXML caches and formulas, without executing or saving the workbook."""

import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "【智慧学习部】FY27预算拆解（1.13亿）-备份.xlsx"
DESTINATION = ROOT / "output/fy27-shadow"


def extract_workbook(source: Path) -> dict:
    """Read saved formulas and caches; do not recalculate, export, or save XLSX."""
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(source, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(source, data_only=True, read_only=False)
    payload = {"source_file": source.name, "sha256": digest,
               "extracted_at": datetime.now(UTC).isoformat(),
               "document_modified_at": str(formulas.properties.modified),
               "source_business_as_of": None, "sheets": {}}
    for sheet in formulas:
        rows = []
        for row in sheet:
            populated = {}
            for cell in row:
                if cell.value is None:
                    continue
                formula = cell.value if cell.data_type == "f" else None
                populated[cell.column_letter] = {
                    "value": cached[sheet.title][cell.coordinate].value,
                    "formula": formula, "format": cell.number_format,
                    "data_type": cached[sheet.title][cell.coordinate].data_type,
                }
            if populated:
                rows.append({"row": row[0].row, "cells": populated})
        payload["sheets"][sheet.title] = {"max_row": sheet.max_row, "max_column": sheet.max_column,
                                         "rows": rows, "merged_ranges": [str(r) for r in sheet.merged_cells.ranges]}
    formulas.close()
    cached.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
        raise ValueError("FY27 Excel 在提取过程中变化，未接受本次提取")
    return payload


def main() -> None:
    # Keep the first extraction immutable. Later source adoption has an explicit
    # equivalence check and separate manifest in adopt_source.py.
    target = DESTINATION / "source-cells.json"
    if target.exists():
        raise ValueError("历史提取已存在，不覆盖；重新核对请使用 adopt_source.py")
    payload = extract_workbook(SOURCE)
    DESTINATION.mkdir(parents=True, exist_ok=True)
    with target.open("x", encoding="utf-8") as output:
        output.write(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n")
    print('Read-only extraction complete; source SHA-256 unchanged:', payload["sha256"])


if __name__ == '__main__':
    main()
